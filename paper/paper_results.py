"""The aggregate results reported in the paper, and their release workbook.

Single source for Figs. 1-3 and Tables I-II: make_figures.py draws from these
values and this module exports them, so the released numbers and the plotted
ones cannot drift apart. Values are transcribed from FIGURE_RECREATION.md,
which records the curated 115-session result files. Percentages throughout.

    python paper/paper_results.py
"""

from __future__ import annotations

from pathlib import Path

# Fig. 1: accuracy by turn-depth bucket, permitted-abstention run.
FIG1 = {
    "depth": [0, 5, 10, 15],
    "scored_turns_per_arm": [570, 573, 575, 169],
    "full_history_C": [84.21, 83.25, 85.74, 78.70],
    "default_3a_1536": [76.67, 64.92, 70.09, 63.31],
    "default_3a_768": [52.98, 46.42, 50.78, 33.73],
}

# Fig. 2: accuracy by compaction/update round, forced-answer runs at 768.
# None marks rounds a policy never reached; they are not plotted as zeros.
FIG2 = {
    "round": list(range(15)),
    "default_3a": [87.13, 46.20, 45.87, 46.72, 36.02, 31.19] + [None] * 9,
    "default_3a_n": [202, 487, 460, 366, 261, 109] + [None] * 9,
    "prompted_3b": [88.61, 42.83, 44.78, 46.01, 33.21, 25.00] + [None] * 9,
    "prompted_3b_n": [202, 481, 469, 363, 262, 108] + [None] * 9,
    "typed_3c": [85.01, 64.91, 71.93, 73.68, 69.30, 65.22, 75.65, 68.42,
                 73.45, 78.57, 74.77, 70.64, 68.22, 57.53, 53.66],
    "typed_3c_n": [387, 114, 114, 114, 114, 115, 115, 114, 113, 112, 111, 109, 107, 73, 41],
}

# Fig. 3: field presence in carried memory, forced-answer runs at 768.
FIG3 = {
    "field": ["entity", "period", "unit", "value", "metric"],
    "default_3a": [24.70, 82.99, 79.91, 20.83, 61.97],
    "prompted_3b": [24.48, 82.72, 72.35, 20.61, 57.97],
    "typed_3c": [92.69, 93.91, 86.18, 28.46, 67.15],
}

# Table I: forced-answer outcomes at the 768-token trigger.
TABLE1 = {
    "policy": ["Full history (3c-run C)", "Default prose (3a)", "Prompted prose (3b)", "Typed facts (3c)"],
    "accuracy": [82.9, 48.3, 46.4, 72.4],
    "answered_wrong": [14.3, 38.8, 36.8, 17.6],
    "recovery": [None, "baseline", "-5%", "70%"],
}

# Table II: potentially dependent turns, conditional on earlier model answers.
TABLE2 = {
    "policy": ["Full history", "Default prose @768", "Prompted prose @768", "Typed facts @768"],
    "earlier_correct": [89.4, 76.8, 78.0, 75.1],
    "earlier_wrong": [33.9, 6.8, 5.2, 17.2],
}

# Further values the text reports, with where they appear.
REPORTED = [
    ("IV.A", "Permitted, full history C, accuracy", 83.89, "n = 1,887"),
    ("IV.A", "Permitted, default prose @1536, accuracy", 69.90, "n = 1,887"),
    ("IV.A", "Permitted, default prose @768, accuracy", 48.60, "n = 1,887"),
    ("IV.A", "369 identical questions, @768 compacted once, accuracy", 45.26, ""),
    ("IV.A", "369 identical questions, @1536 not yet compacted, accuracy", 82.66, ""),
    ("IV.A", "Gap 3a@768 vs full history at round 1, points", 35.1, ""),
    ("IV.A", "Gap 3a@768 vs full history at round 5, points", 53.2, ""),
    ("IV.B", "Permitted @768, correct", 48.6, ""),
    ("IV.B", "Permitted @768, answered-wrong", 27.0, ""),
    ("IV.B", "Permitted @768, abstained", 24.4, ""),
    ("IV.B", "Forced @768, correct", 48.3, ""),
    ("IV.B", "Forced @768, answered-wrong", 38.8, ""),
    ("IV.B", "Forced @768, abstained", 12.8, ""),
    ("IV.B", "Forced full-history reference (own run), correct", 84.3, ""),
    ("IV.B", "Forced full-history reference (own run), answered-wrong", 12.6, ""),
    ("IV.C", "Typed facts @1536, accuracy", 74.8, "one company lacks this arm"),
    ("IV.C", "Sonnet summarizer @768, accuracy (from 46.1)", 48.6, "50 sessions; own ceiling 82.7"),
    ("IV.D", "Scorable turns labelled potentially dependent", 33.2, ""),
    ("IV.E", "Routed multi-agent, full history, accuracy", 53.9, "50 sessions"),
    ("IV.E", "Routed multi-agent, default prose @768, accuracy", 43.1, "50 sessions"),
    ("IV.E", "Routed multi-agent, full history, answered-wrong", 9.0, "50 sessions"),
    ("IV.E", "Routed multi-agent, default prose @768, answered-wrong", 20.2, "50 sessions"),
]

OUT = Path(__file__).resolve().parents[1] / "dataset" / "lost-in-handoff-paper-results.xlsx"

README = [
    "Lost in Handoff: aggregate results reported in the paper.",
    "",
    "These are the numbers behind Figs. 1-3, Tables I-II, and other values quoted in the text, "
    "taken from the curated results of the 115-session experiment runs. Turn-level model outputs "
    "are not released.",
    "",
    "The evaluation inputs (source conversations and turns, chained sessions, and cross-period and "
    "cross-entity pairs) are in convfinqa-evaluation-inputs.xlsx.",
    "",
    "Model: claude-haiku-4-5-20251001. Percentages throughout. Compaction triggers of 768 and 1536 "
    "tokens. 'Permitted' and 'forced' refer to whether the agent may abstain.",
    "",
    "Sheets: fig1_accuracy_by_depth, fig2_accuracy_by_round, fig3_field_presence, "
    "table1_forced_outcomes, table2_dependent_turns, reported_values.",
]


def _columns_to_rows(table: dict) -> tuple[list[str], list[list]]:
    header = list(table)
    return header, [list(row) for row in zip(*table.values())]


def main() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "README"
    for line in README:
        ws.append([line])
    ws["A1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 120

    sheets = {
        "fig1_accuracy_by_depth": FIG1,
        "fig2_accuracy_by_round": FIG2,
        "fig3_field_presence": FIG3,
        "table1_forced_outcomes": TABLE1,
        "table2_dependent_turns": TABLE2,
    }
    for name, table in sheets.items():
        header, rows = _columns_to_rows(table)
        sheet = wb.create_sheet(name)
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)

    sheet = wb.create_sheet("reported_values")
    sheet.append(["section", "quantity", "value", "note"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in REPORTED:
        sheet.append(list(row))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
