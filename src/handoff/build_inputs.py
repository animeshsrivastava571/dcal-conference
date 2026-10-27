"""Build lost-in-handoff-dataset.xlsx: the single released workbook.

Two kinds of sheet, labelled so they cannot be confused:

- inputs: public ConvFinQA source material plus the exact sessions and
  cross-session pairs the experiments used;
- results_*: the aggregate numbers behind the paper's figures and tables.

There are no turn-level model predictions, accuracy or memory columns: the
inputs alone cannot recreate those results, and the turn-level outputs are not
released.

    python -m handoff.build_inputs

The run stops with an error unless every sheet count and the first sampled IDs
match the reference values, which is how incorrect sampling is detected.
"""

from __future__ import annotations

import hashlib
import json
import math
import random
import re
import zipfile
from collections import defaultdict
from itertools import combinations
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from handoff.config import DATA_DIR, PROJECT_ROOT
from handoff.paper_results import FIG1, FIG2, FIG3, REPORTED, TABLE1, TABLE2

ARCHIVE = DATA_DIR / "data.zip"
ARCHIVE_URL = "https://raw.githubusercontent.com/czyssrs/ConvFinQA/main/data.zip"
ARCHIVE_SHA256 = "d764271fae60d81b62e6d58dfc481807ebc8cfbcd633811241723c4a2101072a"
OUT = PROJECT_ROOT / "dataset" / "lost-in-handoff-dataset.xlsx"

SESSION_MIN, SESSION_MAX = 15, 20
N_COMPANIES = 115
N_PERIOD_PAIRS = 50
N_ENTITY_PAIRS = 50
MIN_YEAR_GAP = 5
REL_TOL, ABS_TOL = 0.01, 1e-6

EXPECTED = {
    "records": 3458,
    "session_conversations": 505,
    "source_conversations": 648,
    "source_turns": 2413,
    "session_turns": 1895,
    "cross_period_pairs": 162,
    "cross_entity_pairs": 187,
    "period_eligible": 161,
    "entity_eligible": 185,
    "first_session_ticker": "ETR",
    "first_session_conversation": "Single_ETR/2017/page_35.pdf-3",
    "first_period_pair": ("Double_UNP/2009/page_68.pdf", "Double_UNP/2014/page_16.pdf"),
    "first_entity_pair": ("Double_RL/2008/page_21.pdf", "Single_EXPD/2017/page_30.pdf-1"),
}

_ID_RE = re.compile(r"([A-Z0-9.\-]+)/(\d{4})/")


def parse_id(raw: str) -> tuple[str, int]:
    core = raw.split("_", 1)[1] if "_" in raw.split("/")[0] else raw
    match = _ID_RE.match(core)
    if match is None:
        raise ValueError(f"unparseable id {raw!r}")
    return match.group(1), int(match.group(2))


def load_records() -> list[dict]:
    digest = hashlib.sha256(ARCHIVE.read_bytes()).hexdigest()
    if digest != ARCHIVE_SHA256:
        raise SystemExit(f"{ARCHIVE} SHA-256 is {digest}, expected {ARCHIVE_SHA256}")

    records = []
    with zipfile.ZipFile(ARCHIVE) as zf:
        for split in ("train", "dev"):
            for raw in json.loads(zf.read(f"data/{split}.json")):
                ticker, year = parse_id(raw["id"])
                ann = raw["annotation"]
                records.append(
                    {
                        "id": raw["id"],
                        "split": split,
                        "ticker": ticker,
                        "year": year,
                        "questions": ann["dialogue_break"],
                        "golds": ann["exe_ans_list"],
                        "programs": ann["turn_program"],
                        "pre_text": raw["pre_text"],
                        "table": raw["table"],
                        "post_text": raw["post_text"],
                    }
                )
    return records


def filing_text(record: dict) -> str:
    rows = "\n".join(" | ".join(cell for cell in row) for row in record["table"])
    return "\n\n".join(
        [" ".join(record["pre_text"]).strip(), rows, " ".join(record["post_text"]).strip()]
    ).strip()


def numeric(value) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def build_sessions(by_ticker: dict[str, list[dict]]) -> list[list[dict]]:
    chains: list[list[dict]] = []
    for convs in by_ticker.values():
        chain: list[dict] = []
        for conv in sorted(convs, key=lambda c: (c["year"], c["id"])):
            chain.append(conv)
            total = sum(len(c["questions"]) for c in chain)
            if total > SESSION_MAX:
                chain = [conv]
                continue
            if total >= SESSION_MIN:
                chains.append(chain)
                chain = []

    random.Random(0).shuffle(chains)
    seen: set[str] = set()
    sessions = []
    for chain in chains:
        ticker = chain[0]["ticker"]
        if ticker in seen:
            continue
        seen.add(ticker)
        sessions.append(chain)
        if len(sessions) == N_COMPANIES:
            break
    return sessions


def build_period_pairs(by_ticker: dict[str, list[dict]]) -> list[tuple[dict, dict]]:
    candidates = []
    for convs in by_ticker.values():
        by_year: dict[int, list[dict]] = defaultdict(list)
        for conv in convs:
            by_year[conv["year"]].append(conv)
        for early, late in combinations(sorted(by_year), 2):
            if late - early < MIN_YEAR_GAP:
                continue
            for e in by_year[early]:
                for l in by_year[late]:
                    candidates.append((e, l))

    random.Random(0).shuffle(candidates)
    seen: set[str] = set()
    pairs = []
    for e, l in candidates:
        if e["ticker"] in seen:
            continue
        seen.add(e["ticker"])
        pairs.append((e, l))
        if len(pairs) == N_PERIOD_PAIRS:
            break
    return pairs


def build_entity_pairs(by_ticker: dict[str, list[dict]]) -> list[tuple[dict, dict]]:
    earliest = [min(convs, key=lambda c: (c["year"], c["id"])) for convs in by_ticker.values()]
    random.Random(0).shuffle(earliest)
    pairs = [(earliest[i], earliest[i + 1]) for i in range(0, len(earliest) - 1, 2)]
    return pairs[:N_ENTITY_PAIRS]


def test_turn_rows(reference: dict, test: dict) -> list[dict]:
    """One row per turn of ``test``, flagged against ``reference``'s numeric golds."""
    ref_values = [v for v in (numeric(g) for g in reference["golds"]) if v is not None]
    rows = []
    for i, (q, gold, prog) in enumerate(zip(test["questions"], test["golds"], test["programs"])):
        value = numeric(gold)
        collides = value is not None and any(
            math.isclose(value, r, rel_tol=REL_TOL, abs_tol=ABS_TOL) for r in ref_values
        )
        rows.append(
            {
                "test_turn_index": i,
                "question": q,
                "gold": gold,
                "is_computed": "(" in prog,
                "eligible": value is not None and not collides,
                "ambiguous": collides,
            }
        )
    return rows


def write_sheet(wb: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"


def main() -> None:
    records = load_records()
    by_ticker: dict[str, list[dict]] = {}
    for rec in records:
        by_ticker.setdefault(rec["ticker"], []).append(rec)

    sessions = build_sessions(by_ticker)
    period_pairs = build_period_pairs(by_ticker)
    entity_pairs = build_entity_pairs(by_ticker)

    # Union of every conversation used, in first-appearance order.
    used: dict[str, dict] = {}
    for chain in sessions:
        for conv in chain:
            used.setdefault(conv["id"], conv)
    for a, b in period_pairs + entity_pairs:
        used.setdefault(a["id"], a)
        used.setdefault(b["id"], b)

    source_conversations = [
        [
            c["id"], c["split"], c["ticker"], c["year"],
            json.dumps(c["pre_text"], ensure_ascii=False),
            json.dumps(c["table"], ensure_ascii=False),
            json.dumps(c["post_text"], ensure_ascii=False),
            filing_text(c),
        ]
        for c in used.values()
    ]
    source_turns = [
        [c["id"], i, q, g, p, "(" in p]
        for c in used.values()
        for i, (q, g, p) in enumerate(zip(c["questions"], c["golds"], c["programs"]))
    ]
    session_turns = []
    for chain in sessions:
        ticker = chain[0]["ticker"]
        depth = 0
        for c_index, conv in enumerate(chain):
            for t_index in range(len(conv["questions"])):
                session_turns.append(
                    [f"{ticker}-session", ticker, depth, c_index, conv["id"], t_index, t_index == 0]
                )
                depth += 1

    period_rows = []
    for p_index, (early, late) in enumerate(period_pairs):
        for r in test_turn_rows(early, late):
            period_rows.append(
                [p_index, early["ticker"], early["id"], early["year"], late["id"], late["year"],
                 r["test_turn_index"], r["question"], r["gold"], r["is_computed"],
                 r["eligible"], r["ambiguous"]]
            )
    entity_rows = []
    for p_index, (other, test) in enumerate(entity_pairs):
        for r in test_turn_rows(other, test):
            entity_rows.append(
                [p_index, other["ticker"], other["id"], other["year"],
                 test["ticker"], test["id"], test["year"],
                 r["test_turn_index"], r["question"], r["gold"], r["is_computed"],
                 r["eligible"], r["ambiguous"]]
            )

    longest = max(len(str(v)) for row in source_conversations for v in row)
    if longest > 32_767:
        raise SystemExit(f"a cell holds {longest} characters, over Excel's 32,767 limit")

    observed = {
        "records": len(records),
        "session_conversations": sum(len(c) for c in sessions),
        "source_conversations": len(source_conversations),
        "source_turns": len(source_turns),
        "session_turns": len(session_turns),
        "cross_period_pairs": len(period_rows),
        "cross_entity_pairs": len(entity_rows),
        "period_eligible": sum(r[10] for r in period_rows),
        "entity_eligible": sum(r[11] for r in entity_rows),
        "first_session_ticker": sessions[0][0]["ticker"],
        "first_session_conversation": sessions[0][0]["id"],
        "first_period_pair": (period_pairs[0][0]["id"], period_pairs[0][1]["id"]),
        "first_entity_pair": (entity_pairs[0][0]["id"], entity_pairs[0][1]["id"]),
    }
    failed = {k: (observed[k], v) for k, v in EXPECTED.items() if observed[k] != v}
    for key, expected in EXPECTED.items():
        mark = "ok " if key not in failed else "FAIL"
        print(f"  [{mark}] {key:<28} {observed[key]!s:<70} expected {expected}")
    if failed:
        raise SystemExit("\nSampling does not match the reference; workbook NOT written.")

    readme = [
        ["Lost in Handoff: evaluation inputs and aggregate results. No turn-level predictions."],
        [""],
        ["This workbook has two kinds of sheet. The input sheets hold the public ConvFinQA source "
         "questions, gold answers, programs and filings, and the exact sessions and cross-session pairs "
         "used in the paper. The results_ sheets hold the aggregate numbers behind its figures and "
         "tables. Turn-level model predictions, accuracy and memory contents are not released."],
        [""],
        [f"Source: {ARCHIVE_URL}"],
        [f"SHA-256: {ARCHIVE_SHA256}"],
        ["Splits used: train then dev, in file order (3,458 records). The test split is not used."],
        [""],
        ["Sampling"],
        ["Sessions: per ticker (first-seen order), conversations sorted by (year, id) are chained until "
         "15-20 turns; a chain exceeding 20 restarts from the overflowing conversation. All chains are "
         "shuffled with random.Random(0); the first chain per ticker is kept, up to 115 companies."],
        ["Cross-period pairs: same ticker, years at least 5 apart, every conversation pairing in source "
         "order; shuffled with a new random.Random(0); first pair per ticker, 50 pairs. The later "
         "conversation's turns are the test turns."],
        ["Cross-entity pairs: each ticker's earliest (year, id) conversation; shuffled with a new "
         "random.Random(0); positions 0/1, 2/3, ... form 50 disjoint pairs. The second conversation's "
         "turns are the test turns."],
        ["eligible: numeric gold that does not equal (1% relative, 1e-6 absolute) any numeric gold in "
         "the earlier/other conversation. ambiguous: numeric gold that does."],
        [""],
        ["Input sheets"],
        [f"source_conversations: {len(source_conversations)} conversations (union of session and pair conversations)"],
        [f"source_turns: {len(source_turns)} original turns of those conversations"],
        [f"session_turns: {len(session_turns)} turns across {len(sessions)} chained sessions"],
        [f"cross_period_pairs: {len(period_rows)} test turns from {len(period_pairs)} pairs "
         f"({observed['period_eligible']} eligible)"],
        [f"cross_entity_pairs: {len(entity_rows)} test turns from {len(entity_pairs)} pairs "
         f"({observed['entity_eligible']} eligible)"],
        [""],
        ["Results sheets (aggregate numbers reported in the paper; percentages unless labelled; model "
         "claude-haiku-4-5-20251001; compaction triggers of 768 and 1536 tokens; 'permitted' and "
         "'forced' refer to whether the agent may abstain)"],
        ["results_fig1_by_depth: Fig. 1, accuracy by turn-depth bucket, permitted-abstention run"],
        ["results_fig2_by_round: Fig. 2, accuracy by compaction round, forced-answer runs at 768; "
         "blank cells are rounds a policy never reached"],
        ["results_fig3_fields: Fig. 3, field presence in carried memory, forced-answer runs at 768"],
        ["results_table1: Table I, forced-answer outcomes at the 768-token trigger"],
        ["results_table2: Table II, potentially dependent turns, conditional on earlier model answers"],
        ["results_reported_values: further values quoted in the text, with their section; labels name the unit where a value is a count or a difference in points"],
    ]

    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "README"
    for line in readme:
        ws.append(line)
    ws["A1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 140

    write_sheet(wb, "source_conversations",
                ["conversation_id", "split", "ticker", "filing_year",
                 "pre_text_json", "table_json", "post_text_json", "filing_text_shown"],
                source_conversations)
    write_sheet(wb, "source_turns",
                ["conversation_id", "turn_index", "question", "gold", "program", "is_computed"],
                source_turns)
    write_sheet(wb, "session_turns",
                ["session_id", "ticker", "session_depth", "conversation_index",
                 "conversation_id", "original_turn_index", "starts_conversation"],
                session_turns)
    write_sheet(wb, "cross_period_pairs",
                ["pair_index", "ticker", "earlier_conversation_id", "earlier_year",
                 "later_conversation_id", "later_year", "test_turn_index", "question",
                 "gold", "is_computed", "eligible", "ambiguous"],
                period_rows)
    write_sheet(wb, "cross_entity_pairs",
                ["pair_index", "other_ticker", "other_conversation_id", "other_year",
                 "test_ticker", "test_conversation_id", "test_year", "test_turn_index",
                 "question", "gold", "is_computed", "eligible", "ambiguous"],
                entity_rows)

    for name, table in [
        ("results_fig1_by_depth", FIG1),
        ("results_fig2_by_round", FIG2),
        ("results_fig3_fields", FIG3),
        ("results_table1", TABLE1),
        ("results_table2", TABLE2),
    ]:
        write_sheet(wb, name, list(table), [list(row) for row in zip(*table.values())])
    write_sheet(wb, "results_reported_values",
                ["section", "quantity", "value", "note"], [list(row) for row in REPORTED])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"\nwrote {OUT.relative_to(PROJECT_ROOT)} ({OUT.stat().st_size:,} bytes)")
    print("sheets:", ", ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
